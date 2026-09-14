"""Assemble the subgroup slide deck from the figures and stills."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

DATA = Path("~/Documents/apbs_electrostatics_84_targets").expanduser()
FIG = DATA / "figures"
MOVIES = DATA / "approach_movies"
OUT = DATA / "protein_subgroup_update.pptx"

DARK = RGBColor(0x1A, 0x1A, 0x1A)
GREY = RGBColor(0x59, 0x59, 0x59)
ACCENT = RGBColor(0x21, 0x66, 0xAC)
FONT = "Helvetica Neue"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def textbox(slide, left, top, width, height, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.paragraphs[0].alignment = align
    return frame


def set_run(paragraph, text, size, bold=False, colour=DARK, italic=False):
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = colour
    run.font.name = FONT
    return run


def title_slide(title, subtitle, author):
    slide = prs.slides.add_slide(BLANK)
    frame = textbox(slide, Inches(0.9), Inches(2.4), W - Inches(1.8), Inches(1.2))
    set_run(frame.paragraphs[0], title, 40, bold=True)
    frame = textbox(slide, Inches(0.9), Inches(3.5), W - Inches(1.8), Inches(0.8))
    set_run(frame.paragraphs[0], subtitle, 20, colour=GREY)
    frame = textbox(slide, Inches(0.9), Inches(4.4), W - Inches(1.8), Inches(0.6))
    set_run(frame.paragraphs[0], author, 14, colour=GREY)
    line = slide.shapes.add_shape(1, Inches(0.9), Inches(3.35), Inches(2.2), Emu(18000))
    line.fill.solid(); line.fill.fore_color.rgb = ACCENT
    line.line.fill.background(); line.shadow.inherit = False
    return slide


def content_slide(title):
    slide = prs.slides.add_slide(BLANK)
    frame = textbox(slide, Inches(0.55), Inches(0.32), W - Inches(1.1), Inches(0.7))
    set_run(frame.paragraphs[0], title, 26, bold=True)
    line = slide.shapes.add_shape(1, Inches(0.55), Inches(1.02), W - Inches(1.1), Emu(9000))
    line.fill.solid(); line.fill.fore_color.rgb = RGBColor(0xD5, 0xD5, 0xD5)
    line.line.fill.background(); line.shadow.inherit = False
    return slide


def bullets(slide, items, left, top, width, height, size=15, spacing=10):
    frame = textbox(slide, left, top, width, height)
    first = True
    for item in items:
        indent = 0
        text = item
        if isinstance(item, tuple):
            text, indent = item
        paragraph = frame.paragraphs[0] if first else frame.add_paragraph()
        first = False
        paragraph.level = indent
        paragraph.space_after = Pt(spacing)
        # A blank entry is a spacer, so it must not get a bullet glyph.
        prefix = "" if not text else ("– " if indent else "▪  ")
        set_run(paragraph, prefix + text, size - (1 if indent else 0),
                colour=GREY if indent else DARK)
    return frame


def picture(slide, path, left, top, width=None, height=None):
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def caption(slide, text, left, top, width, size=11):
    frame = textbox(slide, left, top, width, Inches(0.4))
    set_run(frame.paragraphs[0], text, size, colour=GREY, italic=True)


# ---------------------------------------------------------------- 1. title
title_slide("Protein Subgroup Update",
            "Electrostatics of the 84 bound heterodimers: surface potentials, "
            "monomer decomposition, and screened approach curves",
            "Jake Sumner   |   27 August 2026")

# ---------------------------------------------------------------- 2. overview
slide = content_slide("What I did since last time")
bullets(slide, [
    "Set up an APBS/pdb2pqr pipeline and ran it on all 84 bound complexes",
    "Built the same maps for each isolated monomer, on the complex's own grid so "
    "the two can be subtracted directly",
    "Used that to get the interaction potential, complex minus the two chains",
    "Separately, walked each monomer in from 50 Å and computed the Debye-screened "
    "interaction with its partner at every step",
    "Everything is scripted and resumable, so the same code runs on the cluster "
    "for the sampled poses",
], Inches(0.7), Inches(1.5), Inches(7.4), Inches(4.5), size=16, spacing=14)
picture(slide, FIG / "still_apbs_surface.png", Inches(8.3), Inches(1.6), width=Inches(4.4))
caption(slide, "1ugh, APBS surface potential, ±5 kT/e", Inches(8.3), Inches(5.0), Inches(4.4))

# ---------------------------------------------------------------- 3. methods 1
slide = content_slide("Methods: getting the potential")
bullets(slide, [
    "pdb2pqr 3.6.1, PARSE forcefield, pH 7 with PROPKA titration states",
    "APBS 3.4.1, linearised PB on a focused multigrid (mg-auto)",
    ("ε 2.0 solute / 78.54 solvent, 0.150 M 1:1 salt, 298.15 K", 1),
    ("1.4 Å probe, smoothed molecular surface, ~0.5 Å fine grid", 1),
    "Potentials come out in kT/e (1 kT/e = 25.7 mV at 298 K)",
    "Input handling needed real work:",
    ("two files reuse a residue number within a chain (insertion codes were "
     "stripped) — pdb2pqr silently merges them, so I renumber first", 1),
    ("one structure has unresolved side chains that crash the H rebuild; it "
     "falls back to truncating those residues, flagged in the output", 1),
    ("pdb2pqr writes fixed-width PQR, so coordinates collide for structures far "
     "from the origin and APBS rejects the file", 1),
], Inches(0.7), Inches(1.45), Inches(11.9), Inches(5.2), size=15, spacing=8)

# ---------------------------------------------------------------- 4. methods 2
slide = content_slide("Methods: from the grid to per-residue numbers")
bullets(slide, [
    "Solvent-accessible surface sampled Shrake–Rupley, 100 points per atom at "
    "r + 1.4 Å, keeping points outside every other atom",
    "Potential trilinearly interpolated from the grid at each surface point",
    "Residue value is the area-weighted mean over its own points:",
    # phi-bar was written as U+03C6 followed by U+0304 COMBINING MACRON, which
    # PowerPoint does not compose onto the glyph -- the bar renders as a stray
    # dash. Subscript form instead, matching phi_complex / phi_B elsewhere.
    ("φ_res = Σ φₖ aₖ / Σ aₖ,   aₖ = 4π(rᵢ + p)² / M", 1),
    "The weighting matters. PARSE radii run 0.0–2.0 Å, so points differ in area "
    "by ~6×; an unweighted mean is off by up to 4 kT/e on individual residues",
    "Buried residues have no surface points and are NaN, not zero",
    "Join key back to the graphs is aa_id (sequential residue index), not "
    "(chain, residue number), which is not unique in two of these files",
], Inches(0.7), Inches(1.45), Inches(11.9), Inches(5.0), size=15, spacing=11)

# ---------------------------------------------------------------- 5. validation
slide = content_slide("The maps behave the way they should")
picture(slide, FIG / "fig1_residue_validation.png", Inches(0.6), Inches(1.35), width=Inches(8.5))
bullets(slide, [
    "22,689 solvent-exposed residues across the 84",
    "ARG/LYS sit at +1.03 kT/e, ASP/GLU at −1.75",
    "HIS slightly negative, mostly neutral at pH 7",
    "Charge/potential correlation is positive in all 84, minimum 0.11",
    "This is the end-to-end check: a transposed grid axis or a broken residue "
    "mapping would flatten it",
], Inches(9.3), Inches(1.6), Inches(3.6), Inches(4.5), size=13, spacing=10)
caption(slide, "Error bars are SEM.", Inches(0.6), Inches(5.6), Inches(8.5))

# ---------------------------------------------------------------- 6. methods 3
slide = content_slide("Methods: monomers and the interaction potential")
bullets(slide, [
    "Each chain re-solved on its own, but with two things held fixed:",
    ("the complex's exact grid (explicit centre) — APBS would otherwise recentre "
     "on each molecule and offset the lattice by ~0.1 Å", 1),
    ("the complex's charges — re-running PROPKA on an isolated chain can change "
     "titration states, which would put part of the answer into protonation", 1),
    "Then Δφ = φ_complex − Σ φ_chains, voxel by voxel",
    "Because the charge set is identical, the Coulomb part cancels by "
    "superposition. What is left is the change in the dielectric and ionic "
    "boundary on binding, i.e. interface desolvation",
    "Verified the grids are bit-identical for all 168 monomers before subtracting",
], Inches(0.7), Inches(1.45), Inches(7.6), Inches(5.0), size=15, spacing=10)
picture(slide, FIG / "still_interaction_map.png", Inches(8.5), Inches(1.5), width=Inches(4.3))
caption(slide, "1acb, Δφ on the complex surface, ±2 kT/e",
        Inches(8.5), Inches(4.85), Inches(4.3))

# ---------------------------------------------------------------- 7. Δφ result
slide = content_slide("Δφ is confined to the interface")
picture(slide, FIG / "fig3_interaction_localisation.png", Inches(1.5), Inches(1.4), width=Inches(6.2))
bullets(slide, [
    "Median |Δφ| drops by 200–700× going from the contact zone out to 20 Å",
    "That is what a boundary-change effect has to look like",
    "It doubles as a check on the grid alignment: if the lattices were offset, "
    "the signal would be spread over the whole surface instead of sitting at "
    "the seam",
    "Per-residue version is in the CSVs, along with buried SASA",
], Inches(8.1), Inches(1.8), Inches(4.7), Inches(4.2), size=14, spacing=12)

# ---------------------------------------------------------------- 8. methods 4
slide = content_slide("Methods: bringing the monomers together")
bullets(slide, [
    "One chain fixed, the other walked from 50 Å in to the crystal pose in 0.5 Å steps",
    "Screened Coulomb between all atom pairs, using the same PARSE charges:",
    ("U = Σ qᵢ φ_B(rᵢ),   φ_B ∝ Σ qⱼ exp(−r/λ_D)/r", 1),
    ("λ_D = 7.86 Å at 0.150 M, 298.15 K — cytosolic ionic strength, and the same "
     "value the APBS runs used", 1),
    ("summed exactly over all pairs; a cutoff biases the total low", 1),
    "Path is generated outward from the bound pose and reversed, so the last "
    "frame is the crystal structure exactly",
    "Straight pull-out works for 60/84. The rest are steered — slide along the "
    "groove, with small rotations — because no straight line separates them",
    "Frames where the chains still overlap are flagged; those energies are not "
    "physical and should be dropped",
], Inches(0.7), Inches(1.45), Inches(11.9), Inches(5.2), size=15, spacing=8)

# ---------------------------------------------------------------- 9. approach
slide = content_slide("Every interface is electrostatically attractive at contact")
picture(slide, FIG / "fig4_approach_curves.png", Inches(0.6), Inches(1.4), width=Inches(8.6))
bullets(slide, [
    "84/84 attractive at the bound pose",
    "Median −3.6 kcal/mol (−6.1 kT), range −15.4 to ~0",
    "82/84 exceed 1 kT",
    "Curves are flat until roughly a Debye length out, then turn over quickly",
    "Overlapping frames excluded",
], Inches(9.4), Inches(1.7), Inches(3.5), Inches(4.3), size=13, spacing=11)

# ---------------------------------------------------------------- 10. screening
slide = content_slide("Screening, range, and sign reversals")
picture(slide, FIG / "fig5_screening.png", Inches(0.35), Inches(1.3), width=Inches(6.4))
picture(slide, FIG / "fig6_range_and_signflips.png", Inches(6.9), Inches(1.3), width=Inches(6.1))
bullets(slide, [
    "Screening removes ~25% of the interaction at contact (median ratio 0.76) and "
    "essentially all of it beyond ~25 Å",
    "The monopole term alone is nowhere near enough — most of these have small net "
    "charge and the interaction is higher-multipole",
    "|U| reaches 0.5 kT at a median of 11 Å, about 1.4 λ_D",
    "16/84 change sign between long range and contact: net-charge repulsion that "
    "becomes attraction once the complementary patches line up",
], Inches(0.6), Inches(4.75), Inches(12.2), Inches(2.4), size=13, spacing=7)

# ---------------------------------------------------------------- 11. movie
slide = content_slide("Approach movie (1ugh, UNG–UGI)")
poster = MOVIES / "chk_1ugh_far.png"
movie = MOVIES / "1ugh_approach.mp4"
if movie.exists():
    slide.shapes.add_movie(str(movie), Inches(0.8), Inches(1.35),
                           width=Inches(7.7), height=Inches(4.8),
                           poster_frame_image=str(poster), mime_type="video/mp4")
bullets(slide, [
    "Chain I walked in to chain E, 101 frames",
    "Readout shows the separation and the screened interaction energy at that step",
    "Moving chain is coloured by each residue's share of U",
    ("blue stabilising, red destabilising", 1),
    "Same script works for any of the 84; 1kfu is also built",
    "PyMOL session is scripted, so the frames and the numbers on screen come from "
    "the same CSV",
], Inches(8.8), Inches(1.6), Inches(4.0), Inches(4.5), size=13, spacing=10)

# ---------------------------------------------------------------- 12. residues
slide = content_slide("What carries the interaction, and what I'd flag")
picture(slide, FIG / "fig8_residue_contributions.png", Inches(0.7), Inches(1.35), width=Inches(5.1))
bullets(slide, [
    "1ugh at contact: the top contributions are complementary salt bridges "
    "(E/Lys137–I/Glu249, E/Arg195–I/Asp282; file numbering)",
    "A few same-charge pairs push the other way, which is what you'd expect",
    "",
    "Caveats:",
    ("the approach model uses a uniform ε = 78.54, so it has no low-dielectric "
     "interior and no desolvation — contact energies are underestimates", 1),
    ("34/84 have some steric overlap on the path; 1kfu's subunits interdigitate "
     "and cannot be separated rigidly at all", 1),
    ("residue numbering in these files is sequential, not the deposited "
     "numbering, so it can't be cross-referenced directly", 1),
    "",
    "Next: run the pipeline on the sampled poses on the cluster, and see whether "
    "the residue-level terms are useful as GNN edge features",
], Inches(6.2), Inches(1.45), Inches(6.6), Inches(5.4), size=13, spacing=7)

prs.save(str(OUT))
print("wrote", OUT)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
